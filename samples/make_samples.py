"""데모·테스트용 인수인계 샘플 자료 생성기.

가상의 '웹서비스팀 김민수 과장 → 이서연 대리' 인수인계 상황을 PDF / HWPX / DOCX / EML / XLSX / PPTX 로 만든다.
등장하는 기관·인물·업체·연락처는 모두 가상이다.

    python samples/make_samples.py            # samples/ 아래에 생성
    python samples/make_samples.py /tmp/out   # 지정한 폴더에 생성

PDF 생성에는 reportlab 이 필요하다(pip install -e ".[dev]").
"""

from __future__ import annotations

import sys
import zipfile
from email import policy
from email.message import EmailMessage
from email.utils import format_datetime
from datetime import datetime, timedelta, timezone
from pathlib import Path

KST = timezone(timedelta(hours=9))


# --------------------------------------------------------------------------- PDF
def make_pdf(path: Path) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Table, TableStyle

    pdfmetrics.registerFont(UnicodeCIDFont("HYGothic-Medium"))
    body = ParagraphStyle("body", fontName="HYGothic-Medium", fontSize=10.5, leading=16, wordWrap="CJK")
    h1 = ParagraphStyle("h1", parent=body, fontSize=18, leading=26, spaceAfter=10)
    h2 = ParagraphStyle("h2", parent=body, fontSize=14, leading=20, spaceBefore=12, spaceAfter=6)

    def table(rows, widths):
        t = Table([[Paragraph(c, body) for c in r] for r in rows], colWidths=widths)
        t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.6, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke)]))
        return t

    story = [
        Paragraph("웹서비스팀 업무정의서", h1),
        Paragraph("작성: 디지털전략부 웹서비스팀 김민수 과장 / 기준일: 2026년 9월", body),
        Paragraph("1. 업무 개요", h2),
        Paragraph(
            "웹서비스팀은 기관 대표 홈페이지(www.example.or.kr)의 운영 전반을 담당한다. 담당자는 홈페이지 콘텐츠 관리, "
            "웹 접근성 관리, 홈페이지 유지보수 계약 관리, 홈페이지 리뉴얼 사업을 수행한다.",
            body,
        ),
        Paragraph("2. 주요 업무", h2),
        table(
            [
                ["업무명", "주요 내용", "비중"],
                ["홈페이지 콘텐츠 관리", "공지사항·보도자료 게시, 메뉴 구성 변경, 게시물 승인", "40%"],
                ["웹 접근성 관리", "분기별 자체 점검, 연 1회 웹 접근성 품질인증 갱신", "20%"],
                ["유지보수 계약 관리", "유지보수 업체 (주)한빛웹솔루션 관리, 월간 실적 검수 및 대금 지급 요청", "25%"],
                ["홈페이지 리뉴얼", "2026년 하반기 리뉴얼 사업 추진(제안요청서 작성, 입찰)", "15%"],
            ],
            [110, 290, 50],
        ),
        Paragraph("3. 반복 수행 업무", h2),
        Paragraph("가. 매주 월요일 오전: 주간 게시물 현황 점검 및 오래된 공지 정리", body),
        Paragraph("나. 매월 5영업일 이내: 유지보수 월간 실적 검수 후 재무팀에 대금 지급 요청", body),
        Paragraph("다. 분기별(1·4·7·10월): 웹 접근성 자체 점검 및 개선 요청", body),
        Paragraph("라. 매년 11월: 웹 접근성 품질인증 갱신 신청(인증기관 온라인 접수)", body),
        PageBreak(),
        Paragraph("4. 유관 부서", h2),
        Paragraph(
            "재무팀(박지훈 차장)과 유지보수 대금 지급 및 리뉴얼 예산을 협의한다. 정보보호팀과는 웹방화벽 정책을, "
            "IT지원팀과는 서버 및 CMS 계정 발급을 협의한다. 유지보수 업체 담당자와는 장애 대응을 협의한다.",
            body,
        ),
        Paragraph("5. 유의사항", h2),
        Paragraph(
            "SSL 인증서는 매년 12월 중순에 만료되므로 최소 2주 전에 갱신 절차를 시작해야 한다. "
            "홈페이지 게시물은 반드시 부서장 승인 후 게시한다.",
            body,
        ),
    ]
    SimpleDocTemplate(str(path), pagesize=A4, title="웹서비스팀 업무정의서", author="김민수").build(story)
    return path


# --------------------------------------------------------------------------- HWPX
_HP = "http://www.hancom.co.kr/hwpml/2011/paragraph"
_HS = "http://www.hancom.co.kr/hwpml/2011/section"
_HH = "http://www.hancom.co.kr/hwpml/2011/head"


def _hwpx_p(text: str, style: int = 0) -> str:
    from xml.sax.saxutils import escape

    return f'<hp:p paraPrIDRef="0" styleIDRef="{style}"><hp:run charPrIDRef="0"><hp:t>{escape(text)}</hp:t></hp:run></hp:p>'


def _hwpx_table(rows: list[list[str]]) -> str:
    from xml.sax.saxutils import escape

    trs = []
    for r, row in enumerate(rows):
        tcs = []
        for c, cell in enumerate(row):
            tcs.append(
                f'<hp:tc><hp:cellAddr colAddr="{c}" rowAddr="{r}"/><hp:cellSpan colSpan="1" rowSpan="1"/>'
                f'<hp:subList><hp:p paraPrIDRef="0" styleIDRef="0"><hp:run charPrIDRef="0"><hp:t>{escape(cell)}</hp:t>'
                f"</hp:run></hp:p></hp:subList></hp:tc>"
            )
        trs.append("<hp:tr>" + "".join(tcs) + "</hp:tr>")
    return (
        f'<hp:p paraPrIDRef="0" styleIDRef="0"><hp:run charPrIDRef="0"><hp:tbl rowCnt="{len(rows)}" colCnt="{len(rows[0])}">'
        + "".join(trs)
        + "</hp:tbl></hp:run></hp:p>"
    )


def make_hwpx(path: Path) -> Path:
    paragraphs = [
        _hwpx_p("홈페이지 운영 매뉴얼", 1),
        _hwpx_p("1. 사용 시스템 및 접근 권한", 2),
        _hwpx_p("홈페이지 운영에 필요한 시스템과 권한 신청 방법은 아래와 같다."),
        _hwpx_table(
            [
                ["시스템", "용도", "권한 수준", "신청/이관 방법"],
                ["CMS(콘텐츠관리시스템)", "게시물 등록·승인, 메뉴 관리", "관리자", "전자결재 'IT 권한신청서' 제출 → IT지원팀 승인"],
                ["웹방화벽 관리 콘솔", "차단 정책 확인", "조회", "정보보호팀에 메일 요청"],
                ["Google Analytics", "방문 통계 분석", "편집자", "기존 담당자가 계정 권한을 직접 이관"],
                ["도메인·SSL 인증서 관리", "인증서 갱신(만료일 2026-12-15)", "관리자", "총무팀 법인카드로 결제"],
            ]
        ),
        _hwpx_p("2. 장애 대응 절차", 2),
        _hwpx_p("홈페이지 접속 장애가 발생하면 즉시 유지보수 업체 장애 대응 창구로 연락하고, 30분 이내에 IT지원팀과 부서장에게 보고한다."),
        _hwpx_p("장애 원인과 조치 내용은 장애 보고서 양식에 맞추어 다음 영업일까지 작성한다."),
        _hwpx_p("3. 게시물 등록 절차", 2),
        _hwpx_p("부서에서 게시 요청 → 담당자 검토 → 부서장 승인 → CMS 게시 순으로 진행한다. 보도자료는 홍보팀 확인 후 게시한다."),
    ]
    section = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<hs:sec xmlns:hs="{_HS}" xmlns:hp="{_HP}">' + "".join(paragraphs) + "</hs:sec>"
    )
    header = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><hh:head xmlns:hh="{_HH}" version="1.4" secCnt="1">'
        '<hh:refList><hh:styles itemCnt="3">'
        '<hh:style id="0" type="PARA" name="바탕글" engName="Normal" paraPrIDRef="0" charPrIDRef="0" nextStyleIDRef="0"/>'
        '<hh:style id="1" type="PARA" name="개요 1" engName="Outline 1" paraPrIDRef="0" charPrIDRef="0" nextStyleIDRef="0"/>'
        '<hh:style id="2" type="PARA" name="개요 2" engName="Outline 2" paraPrIDRef="0" charPrIDRef="0" nextStyleIDRef="0"/>'
        "</hh:styles></hh:refList></hh:head>"
    )
    content_hpf = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<opf:package xmlns:opf="http://www.idpf.org/2007/opf/" version="" unique-identifier="" id="">'
        "<opf:metadata><opf:title>홈페이지 운영 매뉴얼</opf:title>"
        '<opf:meta name="creator" content="text">김민수</opf:meta></opf:metadata>'
        '<opf:manifest><opf:item id="header" href="Contents/header.xml" media-type="application/xml"/>'
        '<opf:item id="section0" href="Contents/section0.xml" media-type="application/xml"/></opf:manifest>'
        '<opf:spine><opf:itemref idref="header" linear="yes"/><opf:itemref idref="section0" linear="yes"/></opf:spine>'
        "</opf:package>"
    )
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr(zipfile.ZipInfo("mimetype"), "application/hwp+zip")
        zf.writestr("version.xml", '<?xml version="1.0" encoding="UTF-8"?><hv:HCFVersion xmlns:hv="http://www.hancom.co.kr/hwpml/2011/version" major="5" minor="1"/>')
        zf.writestr("Contents/content.hpf", content_hpf)
        zf.writestr("Contents/header.xml", header)
        zf.writestr("Contents/section0.xml", section)
        zf.writestr("Preview/PrvText.txt", "홈페이지 운영 매뉴얼\r\n1. 사용 시스템 및 접근 권한")
    return path


# --------------------------------------------------------------------------- DOCX
def make_docx(path: Path) -> Path:
    import docx

    d = docx.Document()
    d.core_properties.author = "김민수"
    d.core_properties.title = "웹서비스팀 주간회의록"
    d.core_properties.created = datetime(2026, 9, 22, 17, 0)
    d.core_properties.modified = datetime(2026, 9, 22, 17, 30)
    d.add_heading("웹서비스팀 주간회의록 (2026-09-22)", level=1)
    d.add_paragraph("참석: 김민수 과장, 이서연 대리, 최유진 주임 / 장소: 7층 소회의실")
    d.add_heading("1. 진행 중인 과제", level=2)
    t = d.add_table(rows=1, cols=4)
    for cell, text in zip(t.rows[0].cells, ["과제", "진행 현황", "일정", "다음 할 일"]):
        cell.text = text
    for row in [
        ["홈페이지 리뉴얼", "제안요청서(RFP) 초안 80% 작성", "10월 중 입찰 공고 예정", "요구사항 정의서 보완, 예산 확정 확인"],
        ["개인정보처리방침 개정 반영", "법무팀 검토 중", "10/15까지 게시 필요", "법무팀 검토 결과 받아 CMS 게시"],
        ["웹 접근성 개선", "대체텍스트 누락 12건 확인", "10월 말까지 개선", "업체에 수정 요청, 11월 인증 갱신 신청"],
    ]:
        cells = t.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text
    d.add_heading("2. 이슈 및 결정사항", level=2)
    d.add_paragraph("리뉴얼 예산 3억 원은 재무팀 검토 후 이사회 승인을 받아야 확정된다(11월 예정).", style="List Bullet")
    d.add_paragraph("유지보수 업체 장애 대응 시간이 평균 2시간으로 계약 기준(1시간)을 넘어 개선 요청이 필요하다.", style="List Bullet")
    d.add_heading("3. 액션 아이템", level=2)
    d.add_paragraph("김민수: RFP 최종안 10/2까지 공유", style="List Bullet")
    d.add_paragraph("이서연: 웹 접근성 개선 목록 업체 전달", style="List Bullet")
    d.save(str(path))
    return path


# --------------------------------------------------------------------------- EML (+ 첨부)
def make_eml(path: Path) -> Path:
    msg = EmailMessage()
    msg["Subject"] = "[재무팀] 9월 홈페이지 유지보수 대금 지급 관련 요청"
    msg["From"] = "박지훈 <jh.park@example.or.kr>"
    msg["To"] = "김민수 <ms.kim@example.or.kr>"
    msg["Cc"] = "이서연 <sy.lee@example.or.kr>"
    msg["Date"] = format_datetime(datetime(2026, 9, 30, 14, 5, tzinfo=KST))
    msg.set_content(
        "김민수 과장님, 안녕하세요. 재무팀 박지훈입니다.\n\n"
        "9월 홈페이지 유지보수 대금 지급을 위해 월간 실적 검수 결과를 10월 7일까지 보내 주시기 바랍니다.\n"
        "검수표 양식은 첨부한 체크리스트를 참고해 주세요.\n\n"
        "또한 홈페이지 리뉴얼 예산(3억 원)은 11월 이사회 승인 후 확정될 예정이니, 입찰 공고 일정은 그 이후로 잡아 주셔야 합니다.\n\n"
        "감사합니다.\n박지훈 드림 (재무팀 차장, 내선 2345)"
    )
    msg.add_attachment(
        "유지보수 월간 검수 체크리스트\n1. 장애 처리 건수 및 처리 시간 확인\n2. 정기 점검 수행 여부 확인\n3. 요청 작업 완료 여부 확인\n".encode("utf-8"),
        maintype="text",
        subtype="plain",
        filename="검수_체크리스트.txt",
    )
    # 기본 접기(folding)는 한글 제목의 공백을 잃을 수 있어 줄 길이 제한을 넉넉히 둔다.
    path.write_bytes(msg.as_bytes(policy=policy.default.clone(max_line_length=998)))
    return path


# --------------------------------------------------------------------------- XLSX
def make_xlsx(path: Path) -> Path:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "계정 목록"
    ws.append(["시스템", "접속 주소", "계정 유형", "비고"])
    ws.append(["CMS 관리자", "https://cms.example.or.kr", "개인 계정(관리자)", "인수자 명의로 신규 발급 필요"])
    ws.append(["Google Analytics", "analytics.google.com", "공용 계정(편집자)", "비밀번호는 보안USB에 보관"])
    ws.append(["도메인 등록업체", "도메인 관리 콘솔", "법인 계정", "총무팀과 공동 관리"])
    ws2 = wb.create_sheet("연간 일정")
    ws2.append(["월", "업무", "관련 부서"])
    ws2.append(["1·4·7·10월", "웹 접근성 자체 점검", "-"])
    ws2.append(["매월 5영업일", "유지보수 실적 검수 및 대금 지급 요청", "재무팀"])
    ws2.append(["11월", "웹 접근성 품질인증 갱신", "인증기관"])
    ws2.append(["12월 초", "SSL 인증서 갱신", "총무팀"])
    wb.save(str(path))
    return path


# --------------------------------------------------------------------------- PPTX
def make_pptx(path: Path) -> Path:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "2026 웹 접근성 개선 계획"
    slide.placeholders[1].text = "웹서비스팀 / 2026.09"
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "주요 지적사항"
    body = slide.placeholders[1].text_frame
    body.text = "이미지 대체텍스트 누락 12건"
    for line in ["키보드 접근 불가 메뉴 3건", "명도 대비 부족 버튼 5건"]:
        body.add_paragraph().text = line
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "추진 일정"
    body = slide.placeholders[1].text_frame
    body.text = "10월: 유지보수 업체 통해 개선 작업"
    body.add_paragraph().text = "11월: 웹 접근성 품질인증 갱신 신청"
    slide.notes_slide.notes_text_frame.text = "인증 갱신이 늦어지면 마크 사용이 중지되므로 일정 엄수"
    prs.save(str(path))
    return path


def make_all(out_dir: Path) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    return [
        make_pdf(out_dir / "업무정의서_웹서비스팀.pdf"),
        make_hwpx(out_dir / "홈페이지_운영매뉴얼.hwpx"),
        make_docx(out_dir / "주간회의록_2026-09-22.docx"),
        make_eml(out_dir / "재무팀_유지보수대금_요청.eml"),
        make_xlsx(out_dir / "시스템_계정목록.xlsx"),
        make_pptx(out_dir / "웹접근성_개선계획.pptx"),
    ]


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).parent
    for p in make_all(target):
        print("생성:", p)
