"""MS Office 문서 파서: Word(.docx), PowerPoint(.pptx), Excel(.xlsx)."""

from __future__ import annotations

from pathlib import Path

from ..models import Block, ParsedDocument, ParseError
from ..textutil import clean_text, heading_level_from_style, table_to_markdown

MAX_SHEET_ROWS = 5000
TABLE_ROWS_PER_BLOCK = 40


# --------------------------------------------------------------------------- Word
def parse_docx(path: Path) -> ParsedDocument:
    try:
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        document = docx.Document(str(path))
    except Exception as exc:
        raise ParseError(f"Word 문서를 열지 못했습니다(.doc 구버전은 .docx 로 저장 후 올려 주세요): {exc}") from exc

    doc = ParsedDocument(source_name=path.name, source_type="docx")
    props = document.core_properties
    doc.metadata.update(
        {
            "title": clean_text(props.title or ""),
            "author": clean_text(props.author or ""),
            "created_at": props.created.date().isoformat() if props.created else "",
            "modified_at": props.modified.date().isoformat() if props.modified else "",
        }
    )

    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            block = _docx_paragraph(item)
            if block:
                doc.blocks.append(block)
        elif isinstance(item, Table):
            markdown = _docx_table(item)
            if markdown:
                doc.blocks.append(Block(markdown, kind="table"))
    return doc


def _docx_paragraph(paragraph) -> Block | None:
    text = clean_text(paragraph.text)
    if not text:
        return None
    style_name = paragraph.style.name if paragraph.style is not None else ""
    level = heading_level_from_style(style_name)
    if level and len(text) <= 150:
        return Block(text, kind="heading", level=level)
    p_pr = paragraph._p.pPr
    is_list = "list" in (style_name or "").lower() or (p_pr is not None and p_pr.numPr is not None)
    if is_list:
        return Block(f"- {text}", kind="list")
    return Block(text, kind="paragraph")


def _docx_table(table) -> str:
    rows = []
    for row in table.rows:
        cells, seen = [], set()
        for cell in row.cells:
            key = id(cell._tc)
            if key in seen:  # 병합 셀은 같은 객체가 반복된다
                cells.append("")
                continue
            seen.add(key)
            cells.append(cell.text)
        rows.append(cells)
    return table_to_markdown(rows)


# --------------------------------------------------------------------------- PowerPoint
def parse_pptx(path: Path) -> ParsedDocument:
    try:
        from pptx import Presentation

        prs = Presentation(str(path))
    except Exception as exc:
        raise ParseError(f"PowerPoint 문서를 열지 못했습니다(.ppt 구버전은 .pptx 로 저장 후 올려 주세요): {exc}") from exc

    doc = ParsedDocument(source_name=path.name, source_type="pptx")
    props = prs.core_properties
    doc.metadata.update({"title": clean_text(props.title or ""), "author": clean_text(props.author or "")})
    doc.metadata["page_count"] = len(prs.slides)

    for slide_no, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        title = clean_text(title_shape.text_frame.text) if title_shape is not None and title_shape.has_text_frame else ""
        if title:
            doc.blocks.append(Block(title, kind="heading", level=2, page=slide_no))
        image_count = 0
        for shape in _iter_shapes(slide.shapes):
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            if getattr(shape, "has_table", False):
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                markdown = table_to_markdown(rows)
                if markdown:
                    doc.blocks.append(Block(markdown, kind="table", page=slide_no))
            elif getattr(shape, "has_text_frame", False):
                paragraphs = shape.text_frame.paragraphs
                lines = []
                for para in paragraphs:
                    text = clean_text(para.text)
                    if text:
                        bullet = "- " if (para.level or len(paragraphs) > 1) else ""
                        lines.append("  " * para.level + bullet + text)
                if lines:
                    doc.blocks.append(Block("\n".join(lines), kind="paragraph", page=slide_no))
            elif _shape_type_name(shape) == "PICTURE":
                image_count += 1
        if image_count:
            doc.blocks.append(Block(f"[이미지 {image_count}개 포함]", kind="image", page=slide_no))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = clean_text(slide.notes_slide.notes_text_frame.text)
            if notes:
                doc.blocks.append(Block(f"발표자 노트: {notes}", kind="paragraph", page=slide_no))
    return doc


def _shape_type_name(shape) -> str:
    try:
        shape_type = shape.shape_type
    except NotImplementedError:  # 일부 graphicFrame 은 유형을 판별하지 못한다
        return ""
    return getattr(shape_type, "name", "") or ""


def _iter_shapes(shapes):
    for shape in shapes:
        if _shape_type_name(shape) == "GROUP":
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


# --------------------------------------------------------------------------- Excel
def parse_xlsx(path: Path) -> ParsedDocument:
    try:
        import openpyxl

        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    except Exception as exc:
        raise ParseError(f"Excel 문서를 열지 못했습니다(.xls 구버전은 .xlsx 로 저장 후 올려 주세요): {exc}") from exc

    doc = ParsedDocument(source_name=path.name, source_type="xlsx")
    try:
        for sheet_index, ws in enumerate(wb.worksheets, start=1):
            rows: list[list[str]] = []
            truncated = False
            for r_index, row in enumerate(ws.iter_rows(values_only=True)):
                if r_index >= MAX_SHEET_ROWS:
                    truncated = True
                    break
                values = ["" if v is None else str(v) for v in row]
                if any(v.strip() for v in values):
                    rows.append(values)
            if not rows:
                continue
            meta = {"sheet": ws.title}
            doc.blocks.append(Block(f"시트: {ws.title}", kind="heading", level=2, page=sheet_index, meta=meta))
            header, body = rows[0], rows[1:]
            if not body:
                doc.blocks.append(Block(table_to_markdown([header]), kind="table", page=sheet_index, meta=meta))
            for start in range(0, len(body), TABLE_ROWS_PER_BLOCK):
                markdown = table_to_markdown([header] + body[start : start + TABLE_ROWS_PER_BLOCK])
                if markdown:
                    doc.blocks.append(Block(markdown, kind="table", page=sheet_index, meta=meta))
            if truncated:
                doc.warnings.append(f"'{ws.title}' 시트가 너무 커서 앞 {MAX_SHEET_ROWS}행만 반영했습니다.")
    finally:
        wb.close()
    return doc
